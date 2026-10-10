"""BUG-013 in-text citations (Michael Q48 = 1, 2026-10-09: Claude does the work, Michael approves the sheet).
Policy text that cites a section the quote review corrected, or a paragraph that does not exist, is changed to
the citation whose eCFR text supports the sentence. Each edit is an exact string with an expected count; if any
count differs, nothing is written. Sentences that stated the regulation wrongly (066) are corrected to what the
CFR says. Adds an 'In-text changes' sheet to the review workbook.

Usage: python apply_intext_citations.py "<review workbook path>"
"""
import json, os, re, sys

HERE = os.path.dirname(os.path.abspath(__file__))
DATA = os.path.join(os.path.dirname(os.path.dirname(HERE)), "server", "policyTemplates", "data")
NOTICE = re.compile(r"42 CFR 493\.1775(?![\d(])")

# (template, old, new, expected count, reason)
EDITS = [
    ("001_accreditation_body_notification.json", "42 CFR 493.1775(a)", "42 CFR 493.63(a)", 1, "30-day change notice for accredited labs is 493.63(a); 493.1775 is inspection of waiver/PPM labs."),
    ("006_test_ordering.json", "within 30 days per 42 CFR 493.1241(a)", "within 30 days per 42 CFR 493.1241(b)", 1, "Oral requests and the 30-day written authorization are 493.1241(b)."),
    ("016_radiation_safety.json", "obligations at 42 CFR 493.1101(c)", "obligations at 42 CFR 493.1101(d)", 1, "Personnel safety procedures are 493.1101(d)."),
    ("028_lis_validation_verification.json", "CLIA at 42 CFR 493.1281(c) requires verification that test result transmission is accurate", "CLIA at 42 CFR 493.1291(a) requires verification that test result transmission is accurate", 1, "Accurate transmission from data entry to final report is 493.1291(a)."),
    ("028_lis_validation_verification.json", "42 CFR 493.1281(c)", "42 CFR 493.1281(a)", 3, "Method/instrument comparison twice a year is 493.1281(a); (c) is documentation."),
    ("037_corrective_action.json", "42 CFR 493.1282(a)(3)", "42 CFR 493.1291(k)(3)", 2, "493.1282(a)(3) does not exist; keeping the original and corrected report is 493.1291(k)(3)."),
    ("040_microbiology_isolation_id.json", "42 CFR 493.1262(a)", "42 CFR 493.1256(e)(4)", 2, "Media checks are 493.1256(e)(4); 493.1262 is Mycobacteriology."),
    ("066_manual_hematology_qc.json", "have two levels of commercial control material run within each 8 hours of operation and each time a change in reagents occurs, per 42 CFR 493.1261(a).",
     "have two levels of commercial control material run within each 8 hours of operation and each time a change in reagents occurs, which exceeds the 42 CFR 493.1269(a) minimum of one control material each 8 hours with patient specimens and controls tested in duplicate.", 1,
     "Manual cell count QC is 493.1269(a), which requires one control each 8 hours (in duplicate), not two levels; the sentence now says the lab's practice exceeds it."),
    ("066_manual_hematology_qc.json", "Subject to 42 CFR 493.1261(a) two-level, 8-hour control requirement.",
     "Subject to 42 CFR 493.1269(a): one control material each 8 hours of operation, with patient specimens and controls tested in duplicate.", 1,
     "States the 493.1269(a) requirement as written; 493.1261 is Bacteriology."),
    ("073_histopathology_qmp.json", "42 CFR 493.1273(a)(2)", "42 CFR 493.1252(d)", 1, "493.1273(a)(2) does not exist; not using substandard reagents and materials is 493.1252(d)."),
    ("080_virology_qc.json", "42 CFR 493.1256(d)(2)(i)", "42 CFR 493.1256(d)(3)(i)", 1, "Two control concentrations for quantitative procedures is 493.1256(d)(3)(i)."),
    ("098_pretransfusion_testing.json", "per 42 CFR 493.1256 and 42 CFR 493.1273(a)", "per 42 CFR 493.1256(e)(1) and 21 CFR 606.65(c)", 1, "493.1273 is Histopathology. Daily blood bank reagent testing is the 21 CFR 606.65(c) table; lot checks are 493.1256(e)(1)."),
    ("098_pretransfusion_testing.json", "42 CFR 493.1273(a)", "21 CFR 606.65(c)", 2, "Daily reagent QC (anti-human globulin, blood grouping reagents, screening cells: each day of use) is 21 CFR 606.65(c)."),
    ("099_blood_component_handling.json", "42 CFR 493.1273(a)", "21 CFR 606.65(c)", 4, "Same as policy 098."),
    ("100_transfusion_administration.json", "Required by FDA 21 CFR 606.151(d).", "Required by FDA 21 CFR 606.151(e).", 1, "Emergency release documentation signed by a physician is 606.151(e)."),
    ("103_personnel_qualifications.json", "42 CFR 493.1443 (moderate complexity) or 42 CFR 493.1405 (high complexity)", "42 CFR 493.1405 (moderate complexity) or 42 CFR 493.1443 (high complexity)", 1, "The two sections were swapped: 493.1405 is moderate, 493.1443 is high complexity."),
    ("103_personnel_qualifications.json", NOTICE, "42 CFR 493.63(a)", 4, "Director change notice within 30 days is 493.63(a)."),
    ("103_personnel_qualifications.json", "42 CFR 493.1463(b)(5)", "42 CFR 493.1463(b)(4)", 1, "493.1463(b)(5) does not exist; evaluating and documenting testing personnel competency is (b)(4)."),
    ("106_waived_and_point_of_care_testing.json", NOTICE, "42 CFR 493.63(a)", 3, "Location and other changes are notified under 493.63(a)."),
    ("108_health_information_management.json", "42 CFR 493.1281(c) (verification of test transmission accuracy)", "42 CFR 493.1291(a) (accurate transmission of test results to the final report)", 1, "The transmission requirement is 493.1291(a)."),
    ("108_health_information_management.json", "42 CFR 493.1281(c)", "42 CFR 493.1291(a)", 2, "Same as above."),
    ("109_laboratory_governance_and_leadership.json", NOTICE, "42 CFR 493.63(a)", 4, "Notification of changes is 493.63(a)."),
    ("110_infection_prevention_and_standard_precautions.json", "42 CFR 493.1101(c)", "42 CFR 493.1101(d)", 1, "Safety procedures are 493.1101(d)."),
]


# Round 2 (same day): section-level references to a section the quote review moved away from, and sentences
# that attribute a requirement to the wrong section. Found by sweeping every template for the old sections.
EDITS2 = [
    ("001_accreditation_body_notification.json", "Notification is required by 42 CFR 493.1100 and 42 CFR 493.1775 so",
     "Notification is required by 42 CFR 493.51 (certificate of compliance) and 42 CFR 493.63 (certificate of accreditation) so", 1,
     "The 30-day change notice is 493.51 / 493.63; 493.1100 is facility administration and 493.1775 is inspection of waiver/PPM labs."),
    ("004_unsuccessful_pt_response.json", "two of three consecutive events, triggering mandatory cessation of patient testing per 42 CFR 493.821.",
     "two of three consecutive events; unsuccessful participation subjects the laboratory to CMS sanctions under 42 CFR 493.803(b).", 1,
     "493.821 only lists the microbiology PT subspecialties; the consequence of unsuccessful PT is CMS sanctions under 493.803(b), not automatic cessation of testing."),
    ("040_microbiology_isolation_id.json", "Required by 42 CFR 493.1262 (microbiology QC) and 42 CFR 493.1261 (general bacteriology and mycology requirements).",
     "Required by 42 CFR 493.1261 (bacteriology), 42 CFR 493.1262 (mycobacteriology), 42 CFR 493.1263 (mycology), and 42 CFR 493.1264 (parasitology).", 1,
     "Each microbiology subspecialty has its own standard: 493.1261 bacteriology, 1262 mycobacteriology, 1263 mycology, 1264 parasitology."),
    ("066_manual_hematology_qc.json", "42 CFR 493.1261 (hematology), 42 CFR 493.1267 (urinalysis where applicable to body-fluid counts),",
     "42 CFR 493.1269 (hematology),", 1,
     "Hematology is 493.1269 (493.1261 is bacteriology); 493.1267 is routine chemistry, not urinalysis or body-fluid counts."),
    ("078_parasitology.json", "Required by 42 CFR 493.1204 (PT subspecialty parasitology), 42 CFR 493.821 (unsuccessful PT response applies to parasitology), and 42 CFR 493.1262 (microbiology QC).",
     "Required by 42 CFR 493.1204 (Condition: Parasitology), 42 CFR 493.1264 (Standard: Parasitology), and 42 CFR 493.821 (proficiency testing: microbiology includes parasitology).", 1,
     "493.1204 is the parasitology condition and 493.1264 its QC standard; 493.1262 is mycobacteriology."),
    ("098_pretransfusion_testing.json", "The CLIA-required reactivity testing performed at the start of each day", "The required reactivity testing performed at the start of each day", 1,
     "The daily reagent testing is an FDA requirement (21 CFR 606.65(c)), not CLIA."),
    ("099_blood_component_handling.json", "The CLIA-required reactivity testing performed at the start of each day", "The required reactivity testing performed at the start of each day", 1,
     "Same as policy 098."),
    ("099_blood_component_handling.json", "21 CFR 606.65 (storage), 21 CFR 606.122 (labeling), 42 CFR 493.1252 (specimen and reagent integrity), and 42 CFR 493.1273 (blood bank reagent QC).",
     "21 CFR 606.65 (supplies and reagents, including daily reagent testing), 21 CFR 606.121 (container label), 21 CFR 606.122 (circular of information), 42 CFR 493.1252 (specimen and reagent integrity), and 42 CFR 493.1271 (immunohematology, including blood storage).", 1,
     "606.65 is supplies and reagents (not storage); labeling is 606.121 and 606.122 is the circular of information; 493.1273 is histopathology."),
    ("107_molecular_testing.json", ", and 42 CFR 493.1276 (molecular pathology specialty requirements).",
     ". CLIA has no separate molecular pathology standard (42 CFR 493.1276 is clinical cytogenetics), so molecular testing meets the general nonwaived requirements.", 1,
     "493.1276 is Standard: Clinical cytogenetics; CLIA has no molecular pathology standard."),
    ("107_molecular_testing.json", "Molecular genetic testing (germline or somatic) follows the additional requirements of 42 CFR 493.1276:",
     "Molecular genetic testing (germline or somatic) follows the general nonwaived requirements (42 CFR 493.1251, 493.1253 and 493.1256) and this laboratory's added requirements:", 1,
     "Same: the listed elements are lab practice, not 493.1276 requirements."),
    ("107_molecular_testing.json", "subject to additional requirements at 42 CFR 493.1276 beyond general molecular QC.",
     "CLIA has no molecular-specific standard; the general nonwaived requirements and this policy's added requirements apply.", 1,
     "Same."),
    ("111_hct_p_human_cells_tissues.json", "21 CFR 1271.155 (donor screening)", "21 CFR 1271.75 (donor screening)", 1,
     "Donor screening is 1271.75; 1271.155 is exemptions and alternatives."),
    ("111_hct_p_human_cells_tissues.json", "donor eligibility documentation per 21 CFR 1271.155 through 1271.170;", "donor eligibility documentation per 21 CFR 1271.55;", 1,
     "The records that must accompany an HCT/P after the donor-eligibility determination are 1271.55; 1271.155-1271.170 are other current good tissue practice sections."),
]


# Round 3 (same day): in-text citations whose own description contradicts the eCFR section, found by checking
# every "CFR x (description)" against the section heading and text.
EDITS3 = [
    ("068_histocompatibility.json", "Required by 42 CFR 493.1262 (specialty histocompatibility).", "Required by 42 CFR 493.1278 (Standard: Histocompatibility).", 1,
     "Histocompatibility is 493.1278; 493.1262 is mycobacteriology."),
    ("071_radioactive_tissue.json", "42 CFR 493.1101(c) (safety from physical hazards)", "42 CFR 493.1101(d) (safety from physical hazards)", 1, "Safety procedures are 493.1101(d)."),
    ("072_electron_microscope_safety.json", "42 CFR 493.1101(c) (physical hazard safety)", "42 CFR 493.1101(d) (physical hazard safety)", 1, "Safety procedures are 493.1101(d)."),
    ("079_radiobioassay_qc.json", "42 CFR 493.1101(c) (safety)", "42 CFR 493.1101(d) (safety)", 1, "Safety procedures are 493.1101(d)."),
    ("073_histopathology_qmp.json", "42 CFR 493.1249 (preanalytic) / 42 CFR 493.1289 (postanalytic) assessments",
     "42 CFR 493.1249 (preanalytic), 42 CFR 493.1289 (analytic) and 42 CFR 493.1299 (postanalytic) assessments", 1,
     "493.1289 is analytic systems quality assessment; postanalytic is 493.1299."),
    ("103_personnel_qualifications.json", "42 CFR 493.1445 (moderate) or 42 CFR 493.1407 (high)", "42 CFR 493.1407 (moderate) or 42 CFR 493.1445 (high)", 1,
     "Swapped: 493.1407 is moderate and 493.1445 high complexity director responsibilities."),
    ("103_personnel_qualifications.json", "42 CFR 493.1445 (moderate) or 493.1407 (high)", "42 CFR 493.1407 (moderate) or 493.1445 (high)", 1, "Same swap."),
    ("106_waived_and_point_of_care_testing.json", "against the six CLIA elements at 42 CFR 493.1235 (a streamlined six-element form is acceptable).",
     "against the six competency elements CLIA requires for nonwaived testing (42 CFR 493.1413(b)(8) and 493.1451(b)(8)); CLIA does not require them for waived testing, and this laboratory applies them anyway (a streamlined six-element form is acceptable).", 1,
     "493.1235 (Subpart K) covers nonwaived testing and does not list the six elements; they are 493.1413(b)(8) / 493.1451(b)(8). The lab's practice stays; the claim that CLIA requires it for waived testing goes."),
    ("106_waived_and_point_of_care_testing.json", "CLIA-required six-element competency assessment for staff performing waived testing,",
     "Six-element competency assessment (required by CLIA for nonwaived testing; applied here to waived testing as well) for staff performing waived testing,", 1,
     "Same: CLIA does not require six-element competency for waived testing."),
    ("109_laboratory_governance_and_leadership.json", "42 CFR 493.1100 (overall CLIA structure)", "42 CFR 493.1100 (facility administration)", 1,
     "493.1100 is Condition: Facility administration."),
]

# Round 4: one more found by the stricter description sweep (generic words excluded).
EDITS4 = [
    ("010_system_downtime.json", "42 CFR 493.1276 (immunohematology records)", "42 CFR 493.1105(a)(3)(ii) (immunohematology records)", 1,
     "493.1276 is clinical cytogenetics; immunohematology record retention is 493.1105(a)(3)(ii)."),
]


def main():
    global EDITS
    round2 = any(a in sys.argv for a in ("--round2", "--round3", "--round4"))
    if "--round4" in sys.argv:
        sys.argv.remove("--round4")
        EDITS = EDITS4
    if "--round2" in sys.argv:
        sys.argv.remove("--round2")
        EDITS = EDITS2
    if "--round3" in sys.argv:
        sys.argv.remove("--round3")
        EDITS = EDITS3
    report = sys.argv[1] if len(sys.argv) > 1 else None
    by_file = {}
    for e in EDITS:
        by_file.setdefault(e[0], []).append(e)
    staged, log = {}, []
    for f, edits in by_file.items():
        p = os.path.join(DATA, f)
        raw = open(p, encoding="utf-8", newline="").read()
        before = json.loads(raw)
        new = raw
        for _, old, rep, n, why in edits:
            if isinstance(old, re.Pattern):
                count = len(old.findall(new))
                if count != n:
                    raise SystemExit(f"{f}: expected {n} of {old.pattern}, found {count}; nothing written")
                new = old.sub(rep, new)
                log.append((f, old.pattern.replace("\\", "").replace("(?![d(])", ""), rep, n, why))
            else:
                count = new.count(old)
                if count != n:
                    raise SystemExit(f"{f}: expected {n} of {old!r}, found {count}; nothing written")
                new = new.replace(old, rep)
                log.append((f, old, rep, n, why))
        after = json.loads(new)
        if after.get("cfr_text_blocks") != before.get("cfr_text_blocks"):
            raise SystemExit(f"{f}: an in-text edit touched the quote blocks; nothing written")
        staged[p] = new
    for p, new in staged.items():
        open(p, "w", encoding="utf-8", newline="").write(new)
    print(f"in-text edits: {sum(x[3] for x in log)} citations in {len(staged)} templates")
    if report:
        from openpyxl import load_workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        wb = load_workbook(report)
        if round2 and "In-text changes" in wb.sheetnames:
            ws = wb["In-text changes"]
        else:
            if "In-text changes" in wb.sheetnames:
                del wb["In-text changes"]
            ws = wb.create_sheet("In-text changes")
            ws.append(["Template", "Was", "Now", "Times", "Reason", "Michael OK? (Y, or what to change)"])
        for c in ws[1]:
            c.fill, c.font = PatternFill("solid", fgColor="01696F"), Font(bold=True, color="FFFFFF")
        for x in log:
            ws.append([x[0][:-5], x[1], x[2], x[3], x[4], ""])
        for col, w in zip("ABCDEF", [30, 60, 60, 7, 60, 20]):
            ws.column_dimensions[col].width = w
        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.freeze_panes = "B2"
        wb.save(report)
        print(f"sheet 'In-text changes' written to {report}")


if __name__ == "__main__":
    main()
