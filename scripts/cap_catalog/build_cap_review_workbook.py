#!/usr/bin/env python3
"""Build the CAP catalog review workbook (parking lot #80).

Joins the extractor output for two catalog years with the CAP rows currently
loaded in pt_vendor_programs and writes one workbook for a CAP-enrolled
director to verify BEFORE anything is loaded. Each row carries the proposed
action: add (in the catalog, not loaded), keep (both), retire (loaded, in
neither catalog), or review (low-confidence extraction). Nothing here writes
to the product.

Usage:
  python scripts/cap_catalog/build_cap_review_workbook.py cap_programs_2026.csv cap_programs_2027.csv loaded_cap_programs.json out.xlsx
"""
import csv
import json
import sys
from collections import OrderedDict

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

# Catalog discipline label -> the pt_category values the product already uses.
CATEGORY_MAP = {
    "Hematology": "Hematology", "Clinical Microscopy": "Urinalysis", "Coagulation": "Coagulation",
    "General Chemistry and Therapeutic Drug Monitoring": "General Chemistry", "Special Chemistry": "General Chemistry",
    "Endocrinology": "Endocrinology", "Toxicology": "Toxicology / TDM", "Immunology": "Immunology / Serology",
    "Viral Markers": "Immunology / Serology", "Flow Cytometry": "Immunology / Serology",
    "Transfusion Medicine": "Blood Bank / Immunohematology", "Bacteriology": "Microbiology", "Mycology": "Microbiology",
    "Virology": "Microbiology", "Parasitology": "Microbiology", "Mycobacteriology": "Microbiology",
    "Multidiscipline Microbiology": "Microbiology", "Molecular Microbiology": "Microbiology",
    "Quality Cross Check": "Quality Cross Check", "Calibration Verification/Linearity": "Calibration Verification / Linearity",
    "Point-of-Care Testing": "Point of Care", "Biochemical and Molecular Genetics": "Molecular / Genetics",
    "Next-Generation Sequencing": "Molecular / Genetics", "Cytogenetics": "Molecular / Genetics",
}


def load_catalog(path, year):
    out = OrderedDict()
    for r in csv.DictReader(open(path, encoding="utf-8")):
        for code in [c for c in r["codes"].split(", ") if c]:
            if code not in out:
                out[code] = {"name": r["program_name"], "discipline": r["discipline"], "page": r["page"], "confidence": r["confidence"], "year": year}
    return out


def main():
    if len(sys.argv) < 5:
        print(__doc__); return 2
    c26 = load_catalog(sys.argv[1], 2026)
    c27 = load_catalog(sys.argv[2], 2027)
    loaded = {r["program_code"]: r for r in json.load(open(sys.argv[3], encoding="utf-8")) if r.get("program_code")}
    codes = sorted(set(c26) | set(c27) | set(loaded))

    wb = Workbook()
    ws = wb.active
    ws.title = "CAP programs review"
    headers = ["Code", "Program name (catalog)", "Discipline (catalog)", "Proposed pt_category", "In 2026", "In 2027", "Loaded today", "Loaded name", "Loaded category", "Extraction confidence", "Page (2026/2027)", "Proposed action", "Verifier note"]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF", name="Calibri", size=11)
        cell.fill = PatternFill("solid", fgColor="01696F")
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 20
    counts = {"add": 0, "keep": 0, "retire": 0, "review": 0}
    for code in codes:
        a, b, l = c26.get(code), c27.get(code), loaded.get(code)
        src = b or a
        conf = (src or {}).get("confidence", "")
        if l and not a and not b:
            action = "retire (not in either catalog)"
        elif (a or b) and not l:
            action = "add"
        else:
            action = "keep"
        if src and conf != "high":
            action = "review: " + action
        counts["review" if action.startswith("review") else action.split(" ")[0]] += 1
        ws.append([
            code, (src or {}).get("name", ""), (src or {}).get("discipline", ""), CATEGORY_MAP.get((src or {}).get("discipline", ""), ""),
            "yes" if a else "", "yes" if b else "", "yes" if l else "", (l or {}).get("program_name", ""), (l or {}).get("pt_category", ""),
            conf, f"{(a or {}).get('page', '')}/{(b or {}).get('page', '')}", action, "",
        ])
    widths = [10, 46, 34, 28, 8, 8, 11, 40, 24, 12, 14, 26, 30]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="Calibri", size=10, color="28251D")
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        act = row[11].value or ""
        if act.startswith("retire"): row[11].font = Font(name="Calibri", size=10, bold=True, color="A12C7B")
        elif act.startswith("add"): row[11].font = Font(name="Calibri", size=10, bold=True, color="437A22")
        elif act.startswith("review"): row[11].font = Font(name="Calibri", size=10, bold=True, color="964219")

    about = wb.create_sheet("About", 0)
    about["A1"] = "CAP Surveys program catalog review (parking lot #80)"; about["A1"].font = Font(bold=True, size=14, color="01696F")
    lines = [
        "Purpose: verify the CAP program list before it is loaded into VeritaPT's enrollment picker. Nothing is loaded until a CAP-enrolled director has reviewed this sheet.",
        "Source: CAP Surveys catalogs 2026 and 2027 (PDFs supplied by Lisa 2026-10-07), text-extracted and parsed by scripts/cap_catalog/extract_cap_catalog.py.",
        f"Catalog codes found: {len(c26)} in 2026, {len(c27)} in 2027. Loaded in the product today: {len(loaded)} CAP codes (source: production, compiled 2026-10-06 from the CAP online store).",
        f"Proposed actions: add {counts['add']}, keep {counts['keep']}, retire {counts['retire']}, review {counts['review']}.",
        "Retire = loaded today but in neither catalog (for example FH2P). Retired codes are marked inactive, never deleted, so existing enrollments keep displaying.",
        "How to verify: mark any row to change in the Verifier note column (wrong name, wrong category, code that should not load, code missing). Codes are loaded exactly as approved here.",
        "Proposed pt_category comes from the catalog discipline through a fixed map; blanks need a category from the verifier.",
    ]
    for i, t in enumerate(lines, start=3):
        about[f"A{i}"] = t; about[f"A{i}"].alignment = Alignment(wrap_text=True, vertical="top")
    about.column_dimensions["A"].width = 120
    wb.active = 0
    wb.save(sys.argv[4])
    print(f"wrote {sys.argv[4]}: {len(codes)} rows; add {counts['add']}, keep {counts['keep']}, retire {counts['retire']}, review {counts['review']}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
