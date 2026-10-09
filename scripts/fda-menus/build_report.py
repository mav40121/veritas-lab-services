"""BUG-009 (2026-10-09): full VeritaMap instrument menus from the FDA CLIA database.

Step 1 of the pipeline (review only, writes nothing to the library):
  1. NAME MAP: for every library instrument, the FDA test-system names (all
     spellings) that belong to it, found by matching the instrument's model
     tokens against FDA's 15,000+ system names. Each match is labeled:
       exact     - every model token matched (e.g. "cobas c 501" ~ "cobas c501")
       fallback  - matched only after dropping a trailing model number
                   (e.g. "Atellica CH 930" -> FDA "Atellica CH Analyzer"); REVIEW
       none      - no FDA test system found; never filled in
       manual    - a manual method with no instrument test system
  2. DIFF: the menu FDA lists for those systems (every reagent maker; FDA analyte
     name, complexity, specialty) against the library entry today: tests to add,
     tests not in FDA, probable renames, complexity changes, specialty differences.

Source of truth: FDA clia_detail.zip
  https://www.accessdata.fda.gov/premarket/ftparea/clia_detail.zip
Specialty IDs are FDA's own legend (search form, 2026-10-09).

Usage:
  python scripts/fda-menus/build_report.py --fda <clia_detail.txt> --lib client/src/lib/fdaInstrumentData.json \
         --out-json scripts/fda-menus/name_map.json --out-xlsx "<report.xlsx>"
"""
import argparse
import collections
import csv
import datetime
import json
import re

FDA_SPECIALTY = {
    "1": "Urinalysis", "2": "General Chemistry", "3": "General Immunology", "4": "Hematology",
    "5": "Immunohematology", "6": "Endocrinology", "7": "Toxicology", "8": "Bacteriology",
    "9": "Mycobacteriology", "10": "Virology", "11": "Parasitology", "12": "Mycology",
    "13": "Cytology", "14": "Cytogenetics", "15": "Histocompatibility", "16": "Pathology",
    "17": "Syphilis Serology",
}
# Words that name the company, not the analyzer; stripped from the front of a
# library key before matching its model tokens.
VENDOR_WORDS = {
    "abbott", "roche", "siemens", "healthineers", "beckman", "coulter", "ortho", "quidelortho", "quidel",
    "sysmex", "mindray", "werfen", "instrumentation", "laboratory", "il", "biomerieux", "bio", "merieux",
    "stago", "radiometer", "nova", "biomedical", "cepheid", "horiba", "rad", "biorad", "tosoh", "bd",
    "hologic", "cellavision", "immucor", "grifols", "copan", "luminex", "diasorin", "sebia", "agilent",
    "streck", "polymedco", "medtox", "labcorp", "thermo", "fisher", "sciex", "waters", "hemocue", "itc",
    "qiagen", "abaxis", "alcor", "scientific", "advanced", "instruments", "iris", "hemochron",
}


# FDA names that belong to a library analyzer but share none of its model tokens.
# Keys are regexes on the normalized library key; values are regexes on normalized
# FDA system names. Every entry is a family or module name for the same analyzer.
FAMILY_EXTRA = {
    r"\bcobas pure\b": [r"\bcobas c 303\b", r"\bcobas e 402\b", r"\bcobas pure\b"],
    r"\bcobas (8000 c 702|c 702)\b": [r"\bcobas 8000 ise\b"],
    r"\babl 800\b": [r"\babl 8\d\d flex\b"],
    r"\bortho vision max\b": [r"\bortho vision analy[sz]er\b"],
    r"\bimmage 800\b": [r"\bimmage system\b"],
    r"\bi stat 1\b": [r"\bi stat chem 8\b"],
    r"\bdimension vista\b": [r"\bdimension vista (integrated )?system\b"],
}
# A family name: the model tokens minus the trailing model number or suffix, then a
# generic word ("Atellica CI Analyzer", "Dimension Vista System", "IMMULITE 2000 Analyzer").
FAMILY_WORDS = r"(analy[sz]ers?|systems?|integrated|instruments?)"
MODEL_SUFFIXES = {"xpi", "sr", "plus", "max", "t"}
# Libraries that list each reported parameter where FDA lists the test (CBC,
# WBC differential, blood gases, urine sediment). Library parameters are kept.
PARAMETER_LEVEL = {"Hematology", "Coagulation", "Blood Gas", "Urinalysis"}


def family_regex(toks):
    if len(toks) < 2 or not (re.fullmatch(r"\d+", toks[-1]) or toks[-1] in MODEL_SUFFIXES):
        return None
    # The family stem must be specific: two or more words, one of them a real word
    # ("dimension vista", "architect c", "atellica ci"). A lone short stem ("m",
    # "ca", "dm", "advia") matches unrelated systems.
    stem = toks[:-1]
    if len(stem) < 2 or not any(len(t) >= 5 and t.isalpha() for t in stem):
        return None
    return re.compile(r"(?<![a-z0-9])" + r"\s*".join(re.escape(t) for t in toks[:-1]) + r"\s+" + FAMILY_WORDS + r"\b")


def norm(s):
    s = s.lower().replace("ortho-clinical", "ortho clinical").replace("&", " and ")
    s = re.sub(r"[^a-z0-9]+", " ", s)
    s = re.sub(r"(?<=[a-z])(?=\d)|(?<=\d)(?=[a-z])", " ", s)  # c4000 -> c 4000
    return re.sub(r"\s+", " ", s).strip()


def model_tokens(key, vendor):
    # "(c702 module)" names the module FDA lists ("cobas 8000 (c702 module)" ->
    # "cobas c 702"); any other parenthetical ("(POC)", "(Toxicology)",
    # "(XN-10, XN-20)") describes the entry and is dropped for matching.
    m = re.search(r"\(([^)]*)\bmodule\)", key, re.I)
    base = re.sub(r"\([^)]*\)", " ", key)
    toks = norm(base).split()
    if m:
        vt0 = set(norm(vendor).split()) | VENDOR_WORDS
        lead = [t for t in toks if t not in vt0][:1]
        return lead + norm(m.group(1)).split()
    vt = set(norm(vendor).split()) | VENDOR_WORDS
    while toks and toks[0] in vt:
        toks.pop(0)
    return toks


def token_regex(toks):
    return re.compile(r"(?<![a-z0-9])" + r"\s*".join(re.escape(t) for t in toks) + r"(?![a-z0-9])")


def aliases(name):
    """FDA's own aliases for an analyte: the full name, the leading phrase, and each
    parenthetical ("White blood cell count (leukocyte count) (WBC)" -> WBC ...)."""
    out = {loose(name), loose(name.split("(")[0])}
    for p in re.findall(r"\(([^)]*)\)", name):
        out.add(loose(p))
    return {x for x in out if len(x) >= 2}


def loose(a):
    a = re.sub(r"\([^)]*\)", " ", a.lower())
    a = re.sub(r"\b(total|serum|plasma|level|assay|quantitative)\b", " ", a)
    return re.sub(r"[^a-z0-9]+", "", a)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--fda", required=True)
    ap.add_argument("--lib", required=True)
    ap.add_argument("--out-json", required=True)
    ap.add_argument("--out-xlsx", required=True)
    a = ap.parse_args()

    rows = list(csv.DictReader(open(a.fda, encoding="latin-1"), delimiter="|"))
    by_system = collections.defaultdict(list)
    for r in rows:
        by_system[r["TEST_SYSTEM_NAME"].strip()].append(r)
    sys_norm = {s: norm(s) for s in by_system}
    # Records whose qualifier names an analyzer ("cobas e602 analyzer", "AU680 Beckman
    # Coulter Analyzers"); matched on the qualifier as well as the system name.
    qual_recs = []
    for r in rows:
        q = " ".join(x for x in ((r.get("QUALIFIER1") or "").strip(), (r.get("QUALIFIER2") or "").strip()) if x)
        if re.search(r"analy[sz]er|system", q, re.I):
            r["_qn"], r["_q"] = norm(q), q
            qual_recs.append(r)
    lib = json.load(open(a.lib, encoding="utf-8"))

    # Library duplicates (same name apart from case/spacing).
    dup = collections.defaultdict(list)
    for k in lib:
        dup[norm(k)].append(k)
    duplicates = [ks for ks in dup.values() if len(ks) > 1]

    name_map, summary, changes = {}, [], []
    for key in sorted(lib):
        entry = lib[key]
        vendor = entry.get("vendor", "")
        if "manual" in vendor.lower() or "manual" in key.lower():
            name_map[key] = {"match": "manual", "fda_systems": []}
            summary.append([key, vendor, "manual method", len(entry["tests"]), "", 0, 0, "manual", "", "", "", "", "", "Manual method; no instrument test system in FDA. Left as is."])
            continue
        toks = model_tokens(key, vendor)
        match, systems, rx, rx_alt = "none", [], None, None
        if toks:
            rx = token_regex(toks)
            # Roche also writes "cobas 602" for cobas e 602.
            if len(toks) == 3 and toks[0] == "cobas" and len(toks[1]) == 1 and toks[2].isdigit():
                rx_alt = token_regex([toks[0], toks[2]])
            systems = [s for s, n in sys_norm.items() if rx.search(n) or (rx_alt and rx_alt.search(n))]
            if systems:
                match = "exact"
                nk = norm(re.sub(r"\([^)]*\bmodule\)", lambda m: m.group(0)[1:-7], key))
                fam = family_regex(toks)
                extra = [s for s, n in sys_norm.items() if fam and fam.search(n)]
                for kr, pats in FAMILY_EXTRA.items():
                    if re.search(kr, nk):
                        extra += [s for s, n in sys_norm.items() if any(re.search(p_, n) for p_ in pats)]
                extra = [s for s in dict.fromkeys(extra) if s not in systems]
                if extra:
                    systems += extra
                    match = "exact+family"
            elif len(toks) > 1 and re.fullmatch(r"\d+", toks[-1]):
                rx2 = token_regex(toks[:-1])
                systems = [s for s, n in sys_norm.items() if rx2.search(n)]
                if systems:
                    match = "fallback"
        recs = [r for s in systems for r in by_system[s]]
        qual_names = []
        if match.startswith("exact"):
            seen = {id(r) for r in recs}
            qhits = [r for r in qual_recs if id(r) not in seen and (rx.search(r["_qn"]) or (rx_alt and rx_alt.search(r["_qn"])))]
            if qhits:
                recs += qhits
                qual_names = sorted({f'{r["TEST_SYSTEM_NAME"].strip()} [qualifier: {r["_q"]}]' for r in qhits})
        name_map[key] = {"match": match, "model_tokens": toks, "fda_systems": sorted(systems), "qualifier_matches": qual_names, "records": len(recs)}

        fda = {}
        conflicts = {}
        evidence = collections.defaultdict(dict)
        for r in recs:
            an = r["ANALYTE_NAME"].strip()
            cx = r["COMPLEXITY"].strip().upper()
            sp = FDA_SPECIALTY.get(r["SPECIALTY_ID"].strip(), "ID " + r["SPECIALTY_ID"].strip())
            if an in fda and fda[an]["complexity"] != cx:
                conflicts.setdefault(an, set()).update({fda[an]["complexity"], cx})
            fda.setdefault(an, {"complexity": cx, "specialty": sp})
            kit = (r.get("QUALIFIER1") or "").strip() or r["TEST_SYSTEM_NAME"].strip()
            evidence[an].setdefault(f"{cx} {r['DOCUMENT_NUMBER'].strip()} {kit[:60]}", None)
        libt = entry["tests"]
        lib_norm = {re.sub(r"\s+", " ", t.lower()).strip(): t for t in libt}
        fda_norm = {re.sub(r"\s+", " ", t.lower()).strip(): t for t in fda}
        adds = [fda_norm[k] for k in fda_norm if k not in lib_norm]
        removes = [lib_norm[k] for k in lib_norm if k not in fda_norm]
        # Probable renames: the library name equals one of FDA's own aliases for a
        # test FDA lists (full name, leading phrase, or parenthetical abbreviation).
        renames = []
        alias_add = collections.defaultdict(list)
        for x in adds:
            for al in aliases(x):
                alias_add[al].append(x)
        used = set()
        for x in list(removes):
            for cand in alias_add.get(loose(x), []) + alias_add.get(loose(x.split("(")[0]), []):
                if cand not in used:
                    renames.append((x, cand))
                    used.add(cand)
                    break
        ren_l, ren_f = {r[0] for r in renames}, {r[1] for r in renames}
        adds = [x for x in adds if x not in ren_f]
        removes = [x for x in removes if x not in ren_l]
        cx_changes, sp_changes = [], []
        for k, lt in lib_norm.items():
            if k in fda_norm:
                f = fda[fda_norm[k]]
                if fda_norm[k] in conflicts:
                    continue  # reported as "FDA records disagree", never as a change
                if libt[lt]["complexity"].upper() != f["complexity"]:
                    cx_changes.append((lt, libt[lt]["complexity"], f["complexity"]))
                if libt[lt]["specialty"] != f["specialty"]:
                    sp_changes.append((lt, libt[lt]["specialty"], f["specialty"]))
        def ev(an, n=6):
            items = list(evidence.get(an, {}))
            return "; ".join(items[:n]) + (f"; ... +{len(items) - n} more" if len(items) > n else "")
        for x in adds:
            changes.append([key, "add (in FDA, not in library)", x, "", "/".join(sorted(conflicts[x])) if x in conflicts else fda[x]["complexity"], fda[x]["specialty"], ev(x, 3)])
        for x in removes:
            changes.append([key, "not in FDA (library only)", x, libt[x]["complexity"], "", libt[x]["specialty"], ""])
        for old, new in renames:
            changes.append([key, "probable rename", f"{old}  ->  {new}", libt[old]["complexity"], "/".join(sorted(conflicts[new])) if new in conflicts else fda[new]["complexity"], fda[new]["specialty"], ev(new, 3)])
        for t, o, n in cx_changes:
            changes.append([key, "COMPLEXITY CHANGE", t, o, n, libt[t]["specialty"], ev(fda_norm[re.sub(r"\s+", " ", t.lower()).strip()])])
        for t, o, n in sp_changes:
            changes.append([key, "specialty differs", t, libt[t]["complexity"], libt[t]["complexity"], f"{o} -> {n}", ""])
        for an, cxs in conflicts.items():
            lt = lib_norm.get(re.sub(r"\s+", " ", an.lower()).strip())
            changes.append([key, "FDA records disagree", an, libt[lt]["complexity"] if lt else "", "/".join(sorted(cxs)), "REVIEW: depends on the kit or cartridge", ev(an, 10)])
        covered = len(libt) - len(removes)
        coverage = covered / len(libt) if libt else 0
        cat = (entry.get("category") or "")
        if match == "none":
            bucket = "no FDA match"
        elif cat in ("Microbiology", "Molecular"):
            bucket = "review: microbiology/molecular"
        elif match == "fallback":
            bucket = "review: fallback match"
        elif cat in PARAMETER_LEVEL:
            bucket = "review: parameter-level"
        elif re.search(r"\((?![^)]*\bmodule\))[^)]*\)", key, re.I):
            bucket = "review: subset or variant entry"
        elif coverage < 0.5:
            bucket = "review: library menu differs from FDA"
        elif coverage < 0.8:
            bucket = "review: low coverage"
        else:
            bucket = "ready"
        note = {"exact": "", "exact+family": "Includes FDA's family or module names for this analyzer (see Name map).", "fallback": "Matched after dropping the model number; confirm these FDA systems are this analyzer.",
                "none": "No FDA test system found under this name. Left as is until a source is identified."}[match]
        if bucket == "review: microbiology/molecular":
            note = "FDA lists microbiology and molecular systems by organism or panel, not by our panel names; needs its own mapping."
        elif bucket == "review: subset or variant entry":
            note = "The library name carries a qualifier (for example '(Toxicology)'), so this entry may hold only part of the analyzer's menu. FDA's full menu is shown for review; it is not written to this entry."
        elif bucket == "review: parameter-level":
            note = (f"FDA lists the test (CBC, WBC differential, blood gases, urine sediment); the library lists each reported parameter. "
                    f"Library parameters are kept. Adds need a name check so FDA's 'Hemoglobin' is not added beside the library's 'HGB'.")
        elif bucket == "review: library menu differs from FDA":
            note = f"Only {coverage:.0%} of the library's tests appear in FDA's records for this analyzer. The library entry itself may be wrong (tests from a different analyzer)."
        elif bucket == "review: low coverage":
            note = f"FDA's records under these names cover {coverage:.0%} of the library's tests; the analyzer may be listed under other names (combination or family systems)."
        name_map[key].update({
            "bucket": bucket,
            "proposed_adds": [{"analyte": x, "complexity": fda[x]["complexity"], "specialty": fda[x]["specialty"], "evidence": ev(x, 3)}
                              for x in sorted(adds) if x not in conflicts],
            "held_adds_fda_disagrees": [{"analyte": x, "complexity": "/".join(sorted(conflicts[x])), "evidence": ev(x, 10)}
                                        for x in sorted(adds) if x in conflicts],
            "proposed_complexity": [{"analyte": t, "from": o, "to": n, "evidence": ev(fda_norm[re.sub(r"\s+", " ", t.lower()).strip()])}
                                    for t, o, n in cx_changes],
        })
        summary.append([key, vendor, bucket, len(libt), len(fda) if systems else "", len(systems), len(recs), match,
                        len(adds), len(removes), len(renames), len(cx_changes), len(sp_changes), note])

    json.dump({"source": "FDA clia_detail.zip", "built": datetime.date.today().isoformat(), "instruments": name_map},
              open(a.out_json, "w", encoding="utf-8"), indent=1, ensure_ascii=False)

    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    wb = Workbook()
    head = PatternFill("solid", fgColor="01696F")
    def sheet(title, cols, data, widths):
        ws = wb.create_sheet(title)
        ws.append(cols)
        for c in ws[1]:
            c.fill, c.font, c.alignment = head, Font(bold=True, color="FFFFFF"), Alignment(wrap_text=True, vertical="center")
        for row in data:
            ws.append(row)
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[ws.cell(1, i).column_letter].width = w
        ws.freeze_panes = "B2"
        ws.auto_filter.ref = ws.dimensions
        return ws
    about = wb.active
    about.title = "About"
    counts = collections.Counter(s[7] for s in summary)
    buckets = collections.Counter(s[2] for s in summary)
    kinds = collections.Counter(c[1] for c in changes)
    lines = [
        "VeritaMap instrument library vs the FDA CLIA database (BUG-009 review report)",
        f"Built {datetime.date.today().isoformat()} from FDA clia_detail.zip ({len(rows):,} records, {len(by_system):,} test system names). Nothing in the library has been changed.",
        "",
        f"Instruments: {len(lib)}. Name map: exact {counts['exact']}, exact plus family names {counts['exact+family']}, fallback (review) {counts['fallback']}, no FDA match {counts['none']}, manual methods {counts['manual']}.",
        f"Changes found: {kinds['add (in FDA, not in library)']} tests to add, {kinds['not in FDA (library only)']} library tests not in FDA, {kinds['probable rename']} probable renames, {kinds['COMPLEXITY CHANGE']} complexity changes, {kinds['specialty differs']} specialty differences, {kinds['FDA records disagree']} analytes where FDA records disagree.",
        "Readiness: " + ", ".join(f"{k} {v}" for k, v in sorted(buckets.items())) + ". 'ready' = exact name match and FDA covers 80% or more of the library's tests.",
        f"Library duplicates (same name apart from capitalization): {len(duplicates)}" + (": " + "; ".join(" = ".join(d) for d in duplicates) if duplicates else ""),
        "",
        "How to read it",
        "Summary: one row per instrument. 'exact' = every model word matched an FDA test system name; 'fallback' = matched only after dropping the model number, so confirm the FDA systems listed on the Name map sheet are this analyzer; 'none' = no FDA system found, left alone.",
        "Changes: one row per test. COMPLEXITY CHANGE rows are the ones that matter most; 'not in FDA' rows are library tests with no FDA record under these system names; 'probable rename' pairs a library name with FDA's name for the same test.",
        "Specialty differences: the library uses some groupings FDA does not (Blood Gas, Hemostasis, Cardiac, Blood Bank, Microbiology). FDA's own specialty is shown; whether to adopt it is a decision, like electrolytes.",
        "Rules: FDA analyte names, complexity and specialty; every reagent maker run on the platform; nothing is inferred or filled in.",
        "A COMPLEXITY CHANGE is listed only when every FDA record for that test on that analyzer agrees. When records disagree (one kit or cartridge waived, another moderate) the row says 'FDA records disagree' and lists each record; complexity then depends on the kit the lab runs.",
        "FDA evidence column: complexity, FDA document number, and the kit or cartridge (FDA's qualifier) for each record. Search the document number at accessdata.fda.gov to confirm.",
        "Family names: some analyzers are listed by FDA under a family or module name (Atellica CI Analyzer, Dimension Vista System, cobas pure = cobas c 303 + e 402, i-STAT CHEM8+ cartridge). The Name map sheet shows every FDA name used for each instrument.",
        "Proposed write step (after review): add FDA tests to 'ready' instruments and apply agreed complexity changes. Library tests FDA does not list are kept, never deleted; they are parameters, calculated values, or tests under names this report did not pair.",
    ]
    for ln in lines:
        about.append([ln])
    about.column_dimensions["A"].width = 160
    about["A1"].font = Font(bold=True, size=14, color="01696F")
    sheet("Summary", ["Instrument (library)", "Vendor", "Readiness", "Library tests", "FDA tests", "FDA systems", "FDA records", "Match",
                      "Add", "Not in FDA", "Renames", "Complexity changes", "Specialty diffs", "Note"],
          sorted(summary, key=lambda r: (r[2] != "ready", r[2], r[0])), [42, 22, 26, 10, 10, 10, 10, 10, 8, 10, 9, 12, 11, 70])
    sheet("Changes", ["Instrument", "Change", "Test", "Library complexity", "FDA complexity", "Specialty", "FDA evidence (complexity, document, kit)"],
          sorted(changes, key=lambda c: (c[1] != "COMPLEXITY CHANGE", c[1] != "FDA records disagree", c[0], c[1], c[2])), [40, 26, 60, 14, 14, 36, 110])
    nm = [[k, v["match"], len(v.get("fda_systems", [])), v.get("records", 0), " | ".join(v.get("fda_systems", [])),
           " | ".join(v.get("qualifier_matches", []))] for k, v in sorted(name_map.items())]
    sheet("Name map", ["Instrument (library)", "Match", "FDA systems", "FDA records", "FDA test system names",
                       "Also matched by FDA's qualifier field"], nm, [42, 10, 10, 10, 120, 90])
    wb.save(a.out_xlsx)
    print(f"instruments {len(lib)} | match {dict(counts)} | readiness {dict(buckets)} | changes {dict(kinds)} | duplicates {len(duplicates)}")


if __name__ == "__main__":
    main()
